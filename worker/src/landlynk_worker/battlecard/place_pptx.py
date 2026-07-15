"""Render the Place setting pack: a branded deck about the place itself.

Transit (with bus destinations), daily life (shops, schools, health, parks),
eating out and cycling, all factual from OpenStreetMap, plus annual events and
local history (AI-generated, when the story has been added). Complements the
report deck: that one is about the market, this one is about what it is like
to live there. AI slides render only when the story exists and are labelled
for review.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.util import Inches

from .marketing_pptx import _card, _heading_row
from .report_pptx import _GREY, _WHITE, _footer, _hex, _rect, _text


def render_place_pptx(
    record: dict,
    heading: str,
    heading_color: str | None = None,
    logo: bytes | None = None,
    accent: str | None = None,
    map_image: bytes | None = None,
) -> bytes:
    """Render the Place setting pack for one catchment's place record."""
    navy = _hex(heading_color)
    accent_rgb = _hex(accent, "C9A24B")
    theme = {"navy": navy, "accent": accent_rgb}
    story = record.get("story") or {}
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    _cover(prs.slides.add_slide(blank), heading, theme, bool(story))
    builders = [
        lambda s: _transport(s, record, theme, map_image),
        lambda s: _cycling(s, record, theme),
        lambda s: _daily_life(s, record, theme),
        lambda s: _food_and_drink(s, record, theme),
    ]
    if story.get("events"):
        builders.append(lambda s: _events(s, story, theme))
    if story.get("history"):
        builders.append(lambda s: _history(s, story, theme))
    if record.get("civic") or story.get("politics"):
        builders.append(lambda s: _politics(s, record, story, theme))
    for build in builders:
        slide = prs.slides.add_slide(blank)
        build(slide)
        _footer(slide, theme, logo, on_dark=False)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _cover(slide, heading: str, theme: dict, has_story: bool) -> None:  # noqa: ANN001
    _rect(slide, 0, 0, Inches(13.333), Inches(7.5), theme["navy"])
    _text(
        slide,
        Inches(0.8),
        Inches(1.7),
        Inches(11.7),
        Inches(0.5),
        [("PLACE SETTING PACK", 14, theme["accent"], True, False)],
    )
    _text(
        slide,
        Inches(0.8),
        Inches(2.3),
        Inches(11.7),
        Inches(1.4),
        [(heading, 40, _WHITE, True, False)],
        shrink=True,
    )
    _text(
        slide,
        Inches(0.8),
        Inches(3.9),
        Inches(11.0),
        Inches(1.6),
        [
            (
                "Everything around the development: how to get about, where "
                "the shops, schools and green space are, where to eat and "
                "drink"
                + (
                    ", what happens here through the year and how the place "
                    "came to be."
                    if has_story
                    else "."
                ),
                16,
                _WHITE,
                False,
                False,
            )
        ],
        shrink=True,
    )
    _text(
        slide,
        Inches(0.8),
        Inches(6.4),
        Inches(11.7),
        Inches(0.6),
        [
            (
                "Facts from OpenStreetMap, anchored on the development's "
                "postcode. Walk and drive times are approximate, from "
                "straight-line distance."
                + (
                    " Events and history are AI-generated; review before use."
                    if has_story
                    else ""
                ),
                11,
                theme["accent"],
                False,
                True,
            )
        ],
    )


def _fmt_dist(m: int | float | None) -> str:
    if m is None:
        return ""
    m = float(m)
    return f"{m / 1000:.1f} km" if m >= 950 else f"{round(m / 10) * 10} m"


def _station_lines(record: dict, theme: dict) -> tuple[str, list]:
    station = record.get("station")
    if not station:
        return "NEAREST STATION", [
            ("No railway station within 8 km.", 11, _GREY, False, False)
        ]
    kind = "Metro" if station.get("metro") else "Rail"
    lines = [
        (
            f"{kind}  ·  {station['distanceKm']} km  ·  about "
            f"{station['walkMinutes']} min walk or "
            f"{station['driveMinutes']} min drive (approx)",
            12,
            theme["navy"],
            False,
            False,
        )
    ]
    for other in (record.get("otherStations") or [])[:3]:
        label = "Metro" if other.get("metro") else "Rail"
        lines.append(
            (
                f"Also {other['name']} ({label}), {other['distanceKm']} km",
                10,
                _GREY,
                False,
                False,
            )
        )
    return f"NEAREST STATION  ·  {station['name']}", lines


def _transport(slide, record: dict, theme: dict, map_image: bytes | None) -> None:  # noqa: ANN001
    _heading_row(slide, "Transport links", theme)
    title, lines = _station_lines(record, theme)
    _card(slide, Inches(0.6), Inches(1.3), Inches(6.0), Inches(2.4), title, lines, theme)

    services = record.get("busServices") or []
    stops = record.get("busStops") or []
    bus_lines: list = []
    for svc in services[:8]:
        dest = " / ".join(svc.get("destinations") or []) or "local service"
        bus_lines.append(
            (f"{svc['ref']}  to  {dest}", 11, theme["navy"], False, False)
        )
    if not bus_lines and record.get("busRoutes"):
        bus_lines.append(
            (
                "Routes " + "  ·  ".join(record["busRoutes"]),
                12,
                theme["navy"],
                True,
                False,
            )
        )
    if stops:
        nearest = stops[0]
        bus_lines.append(
            (
                f"Nearest stop {nearest['name']}, {_fmt_dist(nearest['distanceM'])}"
                + (
                    f"; also {stops[1]['name']}, {_fmt_dist(stops[1]['distanceM'])}"
                    if len(stops) > 1
                    else ""
                ),
                10,
                _GREY,
                False,
                False,
            )
        )
    bus_lines.append(
        (
            "Bus destinations are the routes' end points; journey times vary "
            "by service.",
            9,
            _GREY,
            False,
            True,
        )
    )
    _card(
        slide,
        Inches(0.6),
        Inches(3.85),
        Inches(6.0),
        Inches(2.8),
        "BUSES FROM THE DOOR",
        bus_lines
        if (services or stops or record.get("busRoutes"))
        else [("No bus stops within 900 m.", 11, _GREY, False, False)],
        theme,
    )

    # The map: development pin (navy) with stations (blue) and stops (red).
    if map_image:
        slide.shapes.add_picture(
            io.BytesIO(map_image),
            Inches(6.8),
            Inches(1.3),
            width=Inches(5.9),
            height=Inches(5.35),
        )
        _text(
            slide,
            Inches(6.8),
            Inches(6.68),
            Inches(5.9),
            Inches(0.3),
            [
                (
                    "Development (navy), stations (blue), bus stops (red). "
                    "Map: OpenStreetMap contributors.",
                    8,
                    _GREY,
                    False,
                    True,
                )
            ],
        )
    else:
        _card(
            slide,
            Inches(6.8),
            Inches(1.3),
            Inches(5.9),
            Inches(2.0),
            "MAP",
            [
                (
                    "Map tiles were not available when this pack was built. "
                    "Re-download to try again.",
                    10,
                    _GREY,
                    False,
                    True,
                )
            ],
            theme,
        )


def _cycling(slide, record: dict, theme: dict) -> None:  # noqa: ANN001
    _heading_row(slide, "Cycling", theme)
    cycles = record.get("cycleRoutes") or []
    if cycles:
        row_h = Inches(4.6 / max(len(cycles[:5]), 2))
        for i, c in enumerate(cycles[:5]):
            title = ("NCN Route " + c["ref"]) if c.get("ref") else "Named route"
            _card(
                slide,
                Inches(0.6),
                Inches(1.3) + i * (row_h + Inches(0.1)),
                Inches(12.1),
                row_h,
                title,
                [
                    (
                        c.get("name")
                        or "Signed national cycle network route within reach.",
                        11,
                        theme["navy"],
                        False,
                        False,
                    )
                ],
                theme,
            )
    else:
        _card(
            slide,
            Inches(0.6),
            Inches(1.3),
            Inches(12.1),
            Inches(1.6),
            "CYCLE ROUTES",
            [
                (
                    "No named or numbered cycle routes within 3 km on "
                    "OpenStreetMap. Local streets and shared paths may still "
                    "offer everyday cycling.",
                    11,
                    _GREY,
                    False,
                    False,
                )
            ],
            theme,
        )
    _text(
        slide,
        Inches(0.6),
        Inches(6.55),
        Inches(12.1),
        Inches(0.4),
        [
            (
                "NCN routes are the National Cycle Network: signed, mapped "
                "long-distance routes that double as everyday commuting spines.",
                10,
                _GREY,
                False,
                True,
            )
        ],
    )


def _politics(slide, record: dict, story: dict, theme: dict) -> None:  # noqa: ANN001
    _heading_row(slide, "The political picture", theme)
    civic = record.get("civic") or {}
    fact_lines: list = []
    if civic.get("constituency"):
        fact_lines.append(
            (f"Constituency: {civic['constituency']}", 12, theme["navy"], False, False)
        )
    if civic.get("council"):
        fact_lines.append(
            (f"Local authority: {civic['council']}", 12, theme["navy"], False, False)
        )
    if civic.get("ward"):
        fact_lines.append(
            (f"Ward: {civic['ward']}", 12, theme["navy"], False, False)
        )
    if civic.get("mp"):
        mp = civic["mp"] + (
            f" ({civic['mpParty']})" if civic.get("mpParty") else ""
        )
        fact_lines.append((f"MP: {mp}", 12, theme["navy"], False, False))
    _card(
        slide,
        Inches(0.6),
        Inches(1.3),
        Inches(5.9),
        Inches(3.2),
        "THE FACTS (ONS AND UK PARLIAMENT)",
        fact_lines
        or [("Civic lookup was unavailable for this pin.", 11, _GREY, False, False)],
        theme,
    )
    politics = story.get("politics") or {}
    pol_lines: list = []
    # Older cached stories carried a model-guessed MP; official records win,
    # so it only shows when the factual lookup has nothing.
    if politics.get("mp") and not civic.get("mp"):
        mp = politics["mp"] + (
            f" ({politics['mpParty']})" if politics.get("mpParty") else ""
        )
        pol_lines.append((f"MP: {mp}", 12, theme["navy"], False, False))
    if politics.get("councilControl"):
        pol_lines.append(
            (
                f"Council control: {politics['councilControl']}",
                12,
                theme["navy"],
                False,
                False,
            )
        )
    if politics.get("commentary"):
        pol_lines.append((politics["commentary"], 10, _GREY, False, False))
    _card(
        slide,
        Inches(6.7),
        Inches(1.3),
        Inches(6.0),
        Inches(3.2),
        "REPRESENTATION AND CHARACTER",
        pol_lines
        or [
            (
                "Add events and history (AI) on the Place profile to complete "
                "this section with the MP, council control and political "
                "character.",
                10,
                _GREY,
                False,
                True,
            )
        ],
        theme,
    )
    if pol_lines:
        _text(
            slide,
            Inches(0.6),
            Inches(6.7),
            Inches(12.1),
            Inches(0.35),
            [
                (
                    "Representation and character are AI-generated as of the "
                    "model's knowledge; verify before use as control can change "
                    "at elections.",
                    10,
                    _GREY,
                    False,
                    True,
                )
            ],
        )


def _list_lines(items: list[dict], theme: dict, limit: int = 6) -> list:
    lines = [
        (
            f"{i['name']}  ·  {_fmt_dist(i['distanceM'])}",
            11,
            theme["navy"],
            False,
            False,
        )
        for i in items[:limit]
    ]
    return lines


def _daily_life(slide, record: dict, theme: dict) -> None:  # noqa: ANN001
    _heading_row(slide, "Daily life within reach", theme)
    quads = [
        ("SHOPS AND ESSENTIALS", (record.get("shops") or []), "No supermarkets or convenience stores mapped within 2 km."),
        ("SCHOOLS", (record.get("schools") or []), "No schools mapped within 2 km on OpenStreetMap."),
        ("HEALTH", (record.get("health") or []), "No GPs, pharmacies or dentists mapped within 2 km."),
        ("PARKS AND LEISURE", (record.get("parks") or []), "No named parks or leisure centres mapped within 1.5 km."),
    ]
    col_w = Inches(5.95)
    row_h = Inches(2.65)
    for i, (title, items, empty) in enumerate(quads):
        col, row = i % 2, i // 2
        _card(
            slide,
            Inches(0.6) + col * (col_w + Inches(0.2)),
            Inches(1.3) + row * (row_h + Inches(0.15)),
            col_w,
            row_h,
            title,
            _list_lines(items, theme)
            or [(empty, 10, _GREY, False, True)],
            theme,
        )


def _food_and_drink(slide, record: dict, theme: dict) -> None:  # noqa: ANN001
    _heading_row(slide, "Food and drink nearby", theme)
    eateries = record.get("restaurants") or []
    if not eateries:
        _card(
            slide,
            Inches(0.6),
            Inches(1.3),
            Inches(12.1),
            Inches(1.5),
            "EATING OUT",
            [
                (
                    "No named restaurants, cafes, pubs or takeaways within "
                    "3 km on OpenStreetMap.",
                    11,
                    _GREY,
                    False,
                    False,
                )
            ],
            theme,
        )
        return
    kinds = {
        "restaurant": "Restaurant",
        "cafe": "Cafe",
        "pub": "Pub",
        "bar": "Bar",
        "fast_food": "Takeaway",
    }
    col_w = Inches(5.95)
    per_col = 6
    for i, e in enumerate(eateries[: per_col * 2]):
        col, row = divmod(i, per_col)
        detail = kinds.get(e.get("type", ""), "Eatery") + (
            f"  ·  {e['cuisine']}" if e.get("cuisine") else ""
        )
        _card(
            slide,
            Inches(0.6) + col * (col_w + Inches(0.2)),
            Inches(1.3) + row * Inches(0.92),
            col_w,
            Inches(0.82),
            f"{e['name']}  ·  {_fmt_dist(e['distanceM'])}",
            [(detail, 9, _GREY, False, False)],
            theme,
        )


def _events(slide, story: dict, theme: dict) -> None:  # noqa: ANN001
    _heading_row(slide, "Events and festivals through the year", theme)
    events = (story.get("events") or [])[:6]
    col_w = Inches(5.95)
    per_col = 3
    for i, e in enumerate(events):
        col, row = divmod(i, per_col)
        lines = []
        if e.get("description"):
            lines.append((e["description"], 10, _GREY, False, False))
        title = e.get("name", "Event")
        if e.get("when"):
            title += f"  ·  {e['when']}"
        _card(
            slide,
            Inches(0.6) + col * (col_w + Inches(0.2)),
            Inches(1.3) + row * Inches(1.7),
            col_w,
            Inches(1.55),
            title,
            lines,
            theme,
        )
    _ai_note(slide, theme)


def _history(slide, story: dict, theme: dict) -> None:  # noqa: ANN001
    _heading_row(slide, "The story of the place", theme)
    _text(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(12.1),
        Inches(4.8),
        [(story.get("history") or "", 15, theme["navy"], False, False)],
        shrink=True,
    )
    _ai_note(slide, theme)


def _ai_note(slide, theme: dict) -> None:  # noqa: ANN001
    _text(
        slide,
        Inches(0.6),
        Inches(6.7),
        Inches(12.1),
        Inches(0.35),
        [
            (
                "AI-generated from the location. Please review before use.",
                10,
                _GREY,
                False,
                True,
            )
        ],
    )
